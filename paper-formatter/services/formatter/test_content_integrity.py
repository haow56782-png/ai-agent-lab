import tempfile
from pathlib import Path

from docx import Document

from docx_formatter import DocxFormatter


def _make_sample_docx(path: Path):
    doc = Document()
    doc.add_paragraph("第一章 绪论")
    doc.add_paragraph("这是一段正文，包含 English words 和  额外空格。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "表格单元格 A"
    table.cell(0, 1).text = "表格单元格 B"
    section = doc.sections[0]
    section.header.paragraphs[0].text = "论文页眉"
    section.footer.paragraphs[0].text = "论文页脚"
    doc.save(path)


def test_content_integrity_matches_after_formatting_only():
    with tempfile.TemporaryDirectory() as temp_dir:
        input_path = Path(temp_dir) / "input.docx"
        output_path = Path(temp_dir) / "output.docx"
        _make_sample_docx(input_path)

        formatter = DocxFormatter(str(input_path))
        formatter.format()
        formatter.save(str(output_path))

        integrity = formatter.get_content_integrity(str(output_path))

        assert integrity["match"] is True
        assert integrity["mismatchCount"] == 0
        assert integrity["originalContentHash"] == integrity["outputContentHash"]


def test_content_integrity_detects_text_mutation():
    with tempfile.TemporaryDirectory() as temp_dir:
        input_path = Path(temp_dir) / "input.docx"
        output_path = Path(temp_dir) / "output.docx"
        _make_sample_docx(input_path)

        formatter = DocxFormatter(str(input_path))
        formatter.format()
        formatter.save(str(output_path))

        mutated_doc = Document(str(output_path))
        mutated_doc.paragraphs[1].text = "这是一段被错误改写的正文。"
        mutated_doc.save(str(output_path))

        integrity = formatter.get_content_integrity(str(output_path))

        assert integrity["match"] is False
        assert integrity["mismatchCount"] >= 1
        assert integrity["mismatches"][0]["original"]["text"] != integrity["mismatches"][0]["output"]["text"]


def test_diff_includes_content_and_package_integrity():
    with tempfile.TemporaryDirectory() as temp_dir:
        input_path = Path(temp_dir) / "input.docx"
        output_path = Path(temp_dir) / "output.docx"
        _make_sample_docx(input_path)

        formatter = DocxFormatter(str(input_path))
        formatter.format()
        formatter.save(str(output_path))

        diff = formatter.get_diff(str(output_path))

        assert diff["integrity"]["contentLevel"]["match"] is True
        assert "packageLevel" in diff["integrity"]


if __name__ == "__main__":
    test_content_integrity_matches_after_formatting_only()
    test_content_integrity_detects_text_mutation()
    test_diff_includes_content_and_package_integrity()
    print("content integrity tests passed")

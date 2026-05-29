#!/usr/bin/env python3
"""parse_enrich 真实单测:构造含公式/三线表/上下标的 docx 并验证增强字段。"""
import copy
import sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import parse_enrich as pe


def _add_omath_paragraph(doc, number_text="(3-1)"):
    """添加一个含 OMML 公式 + 编号的段落。"""
    p = doc.add_paragraph()
    # 构造最小 oMath: m:oMath > m:r > m:t
    m = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    omath = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = "E=mc^2"
    mr.append(mt)
    omath.append(mr)
    p._p.append(omath)
    # 编号 run(普通文本)
    run = p.add_run(number_text)
    return p


def _set_table_borders(table, inside_v="single", inside_h="single"):
    tblPr = table._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for name, val in [("top", "single"), ("bottom", "single"),
                      ("insideH", inside_h), ("insideV", inside_v)]:
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:val"), val)
        borders.append(el)
    tblPr.append(borders)


def _add_sub_sup_run(doc):
    p = doc.add_paragraph()
    r = p.add_run("2")
    rPr = r._element.get_or_add_rPr()
    va = OxmlElement("w:vertAlign")
    va.set(qn("w:val"), "subscript")
    rPr.append(va)
    return p


def _serialize_paragraphs(doc):
    """模拟 extract_paragraphs 的最小序列化。"""
    out = []
    for i, para in enumerate(doc.paragraphs):
        out.append({
            "index": i,
            "text": para.text,
            "runs": [{"text": r.text} for r in para.runs],
        })
    return out


def _serialize_tables(doc):
    return [{"index": ti, "rows": len(t.rows), "cols": len(t.columns)}
            for ti, t in enumerate(doc.tables)]


def test_formula_extraction():
    doc = Document()
    doc.add_paragraph("普通正文段落。")
    _add_omath_paragraph(doc, "(3-1)")
    formulas = pe.extract_formulas(doc)
    assert len(formulas) == 1, f"应识别 1 个公式,实际 {len(formulas)}"
    assert formulas[0]["type"] == "formula"
    assert "(3-1)" in formulas[0]["numbering"]
    assert formulas[0]["is_display"] is True
    print("✓ 公式识别:type=formula, 编号=(3-1), is_display=True")


def test_three_line_table():
    doc = Document()
    # 三线表:无竖线、无内部横线
    t1 = doc.add_table(rows=2, cols=2)
    _set_table_borders(t1, inside_v="none", inside_h="none")
    is3, borders = pe.classify_three_line(t1)
    assert is3 is True, f"应判为三线表,实际 {is3}"
    # 网格表:有竖线
    t2 = doc.add_table(rows=2, cols=2)
    _set_table_borders(t2, inside_v="single", inside_h="single")
    is3b, _ = pe.classify_three_line(t2)
    assert is3b is False, f"应判为非三线表,实际 {is3b}"
    # 无边框信息:不可判定
    t3 = doc.add_table(rows=2, cols=2)
    is3c, _ = pe.classify_three_line(t3)
    assert is3c is None, f"无边框应不可判定,实际 {is3c}"
    print("✓ 三线表:none/none→True, single→False, 无边框→None(不误报)")


def test_sub_sup():
    doc = Document()
    _add_sub_sup_run(doc)
    paras = _serialize_paragraphs(doc)
    pe.enrich_paragraph_runs(doc, paras)
    # 找到含 run 的段落
    target = [p for p in paras if p["runs"]]
    assert target, "应有含 run 的段落"
    assert target[0]["runs"][0]["subscript"] is True
    assert target[0]["runs"][0]["superscript"] is False
    print("✓ 上下标:subscript=True, superscript=False")


def test_full_enrich_no_break():
    """全量增强不破坏原 result 结构,且对空文档健壮。"""
    doc = Document()
    doc.add_paragraph("正文")
    _add_omath_paragraph(doc)
    t = doc.add_table(rows=2, cols=2)
    _set_table_borders(t, inside_v="none", inside_h="none")
    result = {
        "paragraphs": _serialize_paragraphs(doc),
        "tables": _serialize_tables(doc),
        "images": [],
        "structure": [],
    }
    before_keys = set(result.keys())
    out = pe.enrich_parse_result(doc, result)
    # 不删除任何顶层 key
    assert before_keys.issubset(set(out.keys()))
    # 公式进入 structure
    assert any(s.get("type") == "formula" for s in out["structure"])
    # 三线表标记
    assert out["tables"][0]["is_three_line"] is True
    # has_formula 标记存在
    assert any(p.get("has_formula") for p in out["paragraphs"])
    print("✓ 全量增强:结构完整 + 公式入structure + 三线表标记 + has_formula")


def test_empty_doc_robust():
    doc = Document()
    result = {"paragraphs": [], "tables": [], "images": [], "structure": []}
    out = pe.enrich_parse_result(doc, result)
    assert out["structure"] == []
    print("✓ 空文档健壮:无异常,structure 保持空")


if __name__ == "__main__":
    test_formula_extraction()
    test_three_line_table()
    test_sub_sup()
    test_full_enrich_no_break()
    test_empty_doc_robust()
    print("\n全部 5 项测试通过。")

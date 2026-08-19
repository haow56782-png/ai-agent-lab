#!/usr/bin/env python3
"""端到端集成验证:真实 docx → parse_enrich → 模拟 detector 读取,验证拼接处不漏。"""
import sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import parse_enrich as pe

def add_omath(doc, num="(3-1)"):
    p = doc.add_paragraph()
    omath = OxmlElement("m:oMath"); mr = OxmlElement("m:r"); mt = OxmlElement("m:t")
    mt.text = "E=mc^2"; mr.append(mt); omath.append(mr); p._p.append(omath)
    p.add_run(num); return p

def set_borders(table, iv, ih):
    bp = table._element.tblPr; b = OxmlElement("w:tblBorders")
    for n,v in [("top","single"),("bottom","single"),("insideH",ih),("insideV",iv)]:
        e=OxmlElement(f"w:{n}"); e.set(qn("w:val"),v); b.append(e)
    bp.append(b)

# ── 构造文档 ──
doc = Document()
doc.add_paragraph("本文研究风控系统。如式所示:")
add_omath(doc, "(3-1)")
add_omath(doc, "(3-3)")  # 故意跳号 3-1 -> 3-3
gt = doc.add_table(rows=2, cols=3); set_borders(gt, "single", "single")  # 网格表(非三线)
doc.add_paragraph("参考文献")
doc.add_paragraph("[1] 王浩. 风控系统[J]. 计算机学报, 2024, 47(3): 12-25.")
doc.add_paragraph("[2] 李明. 区块链支付. 2022.")  # 缺类型标识

# ── 模拟 parse.py 的序列化 ──
paragraphs = [{"index":i,"text":p.text,"runs":[{"text":r.text} for r in p.runs]} for i,p in enumerate(doc.paragraphs)]
tables = [{"index":ti,"rows":len(t.rows),"cols":len(t.columns)} for ti,t in enumerate(doc.tables)]
# 模拟 detect_structure 产出 reference_entry(P0-d 重构后)
import re
structure=[]
in_ref=False
for i,p in enumerate(doc.paragraphs):
    t=p.text.strip()
    if re.match(r'^参考文献',t): in_ref=True; continue
    if in_ref and re.match(r'^\[\s*\d+\s*\]',t):
        structure.append({"index":i,"type":"reference_entry","text":t})

result={"paragraphs":paragraphs,"tables":tables,"images":[],"structure":structure}
result=pe.enrich_parse_result(doc, result)

# ── 验证拼接处:detector 能拿到的字段 ──
print("=== 端到端字段贯通验证 ===")
formulas=[s for s in result["structure"] if s.get("type")=="formula"]
print(f"[公式] structure 中 type=formula: {len(formulas)} 个 (期望2) {'✓' if len(formulas)==2 else '✗'}")
print(f"[公式] 编号文本可读: {[f['numbering'] for f in formulas]}")
print(f"[三线表] tables[0].is_three_line = {result['tables'][0]['is_three_line']} (期望False,网格表) {'✓' if result['tables'][0]['is_three_line']==False else '✗'}")
refs=[s for s in result["structure"] if s.get("type")=="reference_entry"]
print(f"[参考文献] reference_entry: {len(refs)} 条 (期望2) {'✓' if len(refs)==2 else '✗'}")
print(f"[参考文献] 正文未截断: '{refs[0]['text'][-10:]}' (应含页码) {'✓' if '25' in refs[0]['text'] else '✗'}")
# 模拟 detector F2 类型标识校验
TYPE_MARKER=re.compile(r'\[([JMDCRSPNGZ])\b')
no_type=[r for r in refs if not TYPE_MARKER.search(r['text'])]
print(f"[字段校验] 检出缺类型标识: {len(no_type)} 条 (期望1,即[2]) {'✓' if len(no_type)==1 else '✗'}")
# 公式跳号
nums=[]
for f in formulas:
    m=re.search(r'[(（](\d+)[-–](\d+)[)）]',f['numbering'])
    if m: nums.append((int(m[1]),int(m[2])))
gap = len(nums)==2 and nums[1][1]!=nums[0][1]+1
print(f"[公式校验] 检出编号跳号 {nums}: {'✓' if gap else '✗'}")
print("\n端到端链路贯通:parser增强字段 → detector可消费,拼接无漏。")

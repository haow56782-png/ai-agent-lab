#!/usr/bin/env python3
"""
Tests for the DOCX Structure Parser.

Generates test .docx files and verifies the parser output
against expected structure metadata.

Usage:
    python3 src/test_parse.py
"""

import sys
import json
import os
import tempfile
from pathlib import Path

# Add src to path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from src.parse import parse_docx, detect_structure

try:
    from docx import Document
    from docx.shared import Pt, Cm, Inches, Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("FAIL: python-docx not installed. Run: pip install python-docx")
    sys.exit(1)


PASS = 0
FAIL = 0


def check(name: str, condition: bool, detail: str = ""):
    global PASS, FAIL
    if condition:
        print(f"  ✓ {name}")
        PASS += 1
    else:
        print(f"  ✗ {name}  {detail}")
        FAIL += 1


def make_docx(paragraphs: list) -> str:
    """Create a temporary .docx file from paragraph specs."""
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    doc = Document()

    for spec in paragraphs:
        if spec.get("type") == "heading":
            doc.add_heading(spec["text"], level=spec.get("level", 1))
        elif spec.get("type") == "figure_caption":
            p = doc.add_paragraph()
            run = p.add_run(spec["text"])
            if spec.get("bold"):
                run.bold = True
        elif spec.get("type") == "table_caption":
            doc.add_paragraph(spec["text"]).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif spec.get("type") == "abstract":
            doc.add_paragraph(spec["text"])
        elif spec.get("type") == "cover":
            p = doc.add_paragraph(spec["text"])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(spec.get("font_size", 22))
        elif spec.get("type") == "ref_entry":
            doc.add_paragraph(spec["text"])
        elif spec.get("type") == "toc":
            doc.add_paragraph(spec["text"])
        else:
            doc.add_paragraph(spec.get("text", ""))
            if spec.get("style"):
                doc.paragraphs[-1].style = doc.styles[spec["style"]]

    doc.save(tmp.name)
    return tmp.name


def setup_module():
    """Print test header."""
    print(f"\n{'='*60}")
    print(f" DOCX Parser Test Suite")
    print(f"{'='*60}")


def test_heading_detection():
    print(f"\n--- Heading Detection ---")
    path = make_docx([
        {"type": "heading", "text": "第一章 绪论", "level": 1},
        {"type": "heading", "text": "1.1 研究背景", "level": 2},
        {"type": "body", "text": "研究背景内容"},
        {"type": "heading", "text": "1.2 相关工作", "level": 2},
        {"type": "heading", "text": "第二章 方法", "level": 1},
    ])
    try:
        result = parse_docx(path)
        headings = result["headings"]
        check("Detected 4 headings", len(headings) == 4,
              f"got {len(headings)}")
        if len(headings) >= 2:
            check("H1 correct level", headings[0]["level"] == 1 and headings[0]["text"] == "第一章 绪论")
            check("H2 correct level", headings[1]["level"] == 2 and headings[1]["text"] == "1.1 研究背景")
        check("Structure includes 4 heading entries",
              len([s for s in result["structure"] if s["type"] == "heading"]) == 4)
    finally:
        os.unlink(path)


def test_figure_caption_detection():
    print(f"\n--- Figure Caption Detection ---")
    path = make_docx([
        {"type": "heading", "text": "第三章 实验", "level": 1},
        {"type": "body", "text": "实验设置如下"},
        {"type": "figure_caption", "text": "图3-1 实验结果对比", "bold": True},
        {"type": "body", "text": "从图中可以看出"},
        {"type": "figure_caption", "text": "图3-2 误差曲线", "bold": True},
    ])
    try:
        result = parse_docx(path)
        figs = [s for s in result["structure"] if s["type"] == "figure_caption"]
        check("Detected 2 figure captions", len(figs) == 2,
              f"got {len(figs)}")
        if figs:
            check("Figure caption confidence >= 0.80",
                  all(f["confidence"] >= 0.80 for f in figs))
    finally:
        os.unlink(path)


def test_cover_detection():
    print(f"\n--- Cover Detection ---")
    path = make_docx([
        {"type": "cover", "text": "基于深度学习的图像超分辨率重建研究", "font_size": 26},
        {"type": "cover", "text": "李小明", "font_size": 16},
        {"type": "body", "text": "这是一篇学位论文"},
        {"type": "abstract", "text": "摘要"},
    ])
    try:
        result = parse_docx(path)
        covers = [s for s in result["structure"] if s["type"] == "cover"]
        check("Detected cover elements", len(covers) >= 2,
              f"got {len(covers)}")
        check("Structure has cover before abstract",
              result["structure"][0]["type"] == "cover")
    finally:
        os.unlink(path)


def test_abstract_detection():
    print(f"\n--- Abstract Detection ---")
    path = make_docx([
        {"type": "cover", "text": "论文标题", "font_size": 22},
        {"type": "abstract", "text": "摘要"},
        {"type": "body", "text": "本文研究深度学习在图像超分辨率领域的应用。" * 5},
    ])
    try:
        result = parse_docx(path)
        abstracts = [s for s in result["structure"] if s["type"] == "abstract"]
        check("Detected abstract header", len(abstracts) >= 1,
              f"got {len(abstracts)}")
    finally:
        os.unlink(path)


def test_reference_detection():
    print(f"\n--- Reference Detection ---")
    path = make_docx([
        {"type": "heading", "text": "参考文献", "level": 1},
        {"type": "ref_entry", "text": "[1] Goodfellow I, Bengio Y, Courville A. Deep Learning. MIT Press, 2016."},
        {"type": "ref_entry", "text": "[2] He K, Zhang X, Ren S, et al. Deep Residual Learning for Image Recognition. CVPR, 2016."},
    ])
    try:
        result = parse_docx(path)
        refs = [s for s in result["structure"] if s["type"] == "references_header" or s["type"] == "reference_entry"]
        check("Detected references section", len(refs) >= 3,
              f"got {len(refs)}")
        check("Reference entries parsed",
              len([s for s in result["structure"] if s["type"] == "reference_entry"]) == 2)
    finally:
        os.unlink(path)


def test_table_caption():
    print(f"\n--- Table Caption Detection ---")
    path = make_docx([
        {"type": "table_caption", "text": "表3-1 实验参数设置"},
    ])
    try:
        result = parse_docx(path)
        tables = [s for s in result["structure"] if s["type"] == "table_caption"]
        check("Detected table caption", len(tables) >= 1)
    finally:
        os.unlink(path)


def test_section_margins():
    print(f"\n--- Section Properties ---")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    doc.add_paragraph("正文")
    path = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    doc.save(path)
    try:
        result = parse_docx(path)
        sections = result["sections"]
        check("Section margins extracted", len(sections) >= 1)
        if sections:
            s = sections[0]
            check("Top margin 25mm", 24.5 <= s["margin_top_mm"] <= 26.0,
                  f"got {s['margin_top_mm']}")
            check("Left margin 30mm", 29.0 <= s["margin_left_mm"] <= 31.0,
                  f"got {s['margin_left_mm']}")
    finally:
        os.unlink(path)


def test_empty_document():
    print(f"\n--- Edge Cases ---")
    path = tempfile.NamedTemporaryFile(suffix=".docx", delete=False).name
    Document().save(path)
    try:
        result = parse_docx(path)
        check("Empty doc: sections > 0", len(result["sections"]) >= 1)
        check("Empty doc: metadata ok", result["metadata"]["paragraphs"] >= 0)
        check("Empty doc: no structure", len(result["structure"]) == 0)
    finally:
        os.unlink(path)


def test_metadata():
    print(f"\n--- Metadata ---")
    path = make_docx([
        {"type": "body", "text": "Hello world"},
    ])
    try:
        result = parse_docx(path)
        meta = result["metadata"]
        check("Paragraph count > 0", meta["paragraphs"] >= 1,
              f"got {meta['paragraphs']}")
        check("Sections count > 0", meta["sections"] >= 1)
        check("Tables count >= 0", meta["tables"] >= 0)
    finally:
        os.unlink(path)


def main():
    setup_module()
    test_heading_detection()
    test_figure_caption_detection()
    test_cover_detection()
    test_abstract_detection()
    test_reference_detection()
    test_table_caption()
    test_section_margins()
    test_empty_document()
    test_metadata()

    total = PASS + FAIL
    print(f"\n{'='*60}")
    print(f" Results: {PASS}/{total} passed", end="")
    if FAIL > 0:
        print(f", {FAIL} failed", end="")
    print()
    print(f"{'='*60}\n")
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
